"""Generate synthetic sample resumes (fictional people) in several formats and languages.

    uv run python scripts/make_samples.py

Covers every extraction path: PDF text layer (with hidden hyperlinks), a scanned PDF and a
PNG photo (OCR), DOCX in English and Hindi, and plain text in French.
"""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fpdf import FPDF
from PIL import Image, ImageDraw, ImageFont

SAMPLES = Path(__file__).resolve().parent.parent / "samples"

JOB_DESCRIPTION = """Android Developer (Kotlin)

We are hiring an Android engineer to build and maintain our consumer mobile apps.

Requirements:
- 2+ years of Android development with Kotlin
- Jetpack Compose and MVVM architecture
- Coroutines, Room and Retrofit
- Firebase (Auth, Firestore, Cloud Messaging)
- Publishing apps on the Google Play Store and writing unit tests

Nice to have:
- CI/CD with GitHub Actions and Hilt dependency injection

Benefits:
- Health insurance, learning budget and a competitive salary
"""

PRIYA = """PRIYA SHARMA
Android Developer | Pune, India | priya.sharma@example.com

SUMMARY
Android engineer with 3 years of experience building consumer apps in Kotlin.

EXPERIENCE
Mobile Engineer, ShopKart (2022 - 2025)
- Rebuilt the checkout flow in Jetpack Compose using an MVVM architecture.
- Kotlin Coroutines and Flow for async work; Room for offline caching; Retrofit for REST.
- Integrated Firebase Auth, Firestore and Cloud Messaging for order notifications.
- Published 4 apps on the Google Play Store (1M+ downloads); 80% unit test coverage.
- Set up CI/CD with GitHub Actions; migrated dependency injection to Hilt.

EDUCATION
B.Tech in Computer Engineering, 2022
"""

ARJUN = """ARJUN PATIL - ANDROID DEVELOPER
arjun.patil@example.com | Nagpur

Two years building Android apps with Kotlin and Java.
Built a food delivery app with Jetpack Compose and MVVM.
Used Retrofit, Room and Kotlin Coroutines.
Firebase Authentication and Firestore for user data.
Published 2 apps on the Google Play Store.
"""

KARAN = """KARAN MEHTA - iOS DEVELOPER
karan.mehta@example.com

Four years with Swift, SwiftUI, Combine and Core Data.
Published 3 apps on the App Store.
Unit tests with XCTest; CI with Fastlane.
"""

ANITA = """ANITA DESAI
Chartered Accountant | anita.desai@example.com

Six years of experience in financial accounting, GST filing, payroll processing,
Tally ERP and Excel reporting. Managed statutory audits for manufacturing clients.

EDUCATION
B.Com, Chartered Accountant (CA)
"""

PIERRE = """Pierre Martin - Développeur Android
Lyon, France - pierre.martin@example.com

Trois ans d'expérience en développement d'applications Android avec Kotlin.
Interfaces en Jetpack Compose, architecture MVVM, Coroutines, Room et Retrofit.
Intégration de Firebase (Auth, Firestore, Cloud Messaging).
A publié quatre applications sur le Google Play Store.
Tests unitaires avec JUnit, intégration continue avec GitHub Actions, injection avec Hilt.
"""

SNEHA = [
    "स्नेहा कुलकर्णी - एंड्रॉइड डेवलपर",
    "पुणे | sneha.kulkarni@example.com",
    "सारांश: कोटलिन (Kotlin) में तीन वर्षों का एंड्रॉइड ऐप विकास अनुभव।",
    "अनुभव: Jetpack Compose और MVVM आर्किटेक्चर के साथ शॉपिंग ऐप बनाया।",
    "Coroutines, Room और Retrofit का उपयोग किया। Firebase Auth और Firestore जोड़ा।",
    "Google Play Store पर तीन ऐप प्रकाशित किए। JUnit से यूनिट टेस्ट लिखे।",
]

RAHUL = [
    "RAHUL VERMA - Full-stack Web Developer",
    "rahul.verma@example.com",
    "Two years building web applications with React, Node.js and MongoDB.",
    "Designed REST APIs with Express and deployed them on Firebase Hosting.",
    "Some Java from university projects.",
]


def text_pdf(path: Path, text: str, links: dict[str, str]) -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    pdf.multi_cell(0, 6, text)
    pdf.ln(2)
    for label, url in links.items():  # visible text only says "GitHub" etc.
        pdf.cell(30, 6, label, link=url)
    pdf.output(str(path))


def render(text: str, size: tuple[int, int] = (1240, 900)) -> Image.Image:
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=30)
    y = 60
    for line in text.splitlines():
        draw.text((70, y), line, fill="black", font=font)
        y += 46
    return image


def scanned_pdf(path: Path, text: str) -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.image(render(text), x=10, y=10, w=190)  # an image only: no text layer, needs OCR
    pdf.output(str(path))


def add_hyperlink(paragraph, url: str, text: str) -> None:  # type: ignore[no-untyped-def]
    rel_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link, run, node = OxmlElement("w:hyperlink"), OxmlElement("w:r"), OxmlElement("w:t")
    link.set(qn("r:id"), rel_id)
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)


def docx_file(path: Path, lines: list[str], link: tuple[str, str] | None = None) -> None:
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    if link:
        add_hyperlink(document.add_paragraph("Profile: "), *link)
    buffer = io.BytesIO()
    document.save(buffer)
    path.write_bytes(buffer.getvalue())


def main() -> None:
    SAMPLES.mkdir(exist_ok=True)
    (SAMPLES / "job_description.txt").write_text(JOB_DESCRIPTION, encoding="utf-8")
    text_pdf(
        SAMPLES / "priya_sharma_android.pdf",
        PRIYA,
        # octocat is GitHub's official demo account, so the profile bonus shows real data
        # without looking up a real person.
        {"GitHub": "https://github.com/octocat"},
    )
    scanned_pdf(SAMPLES / "arjun_patil_android_scan.pdf", ARJUN)
    render(KARAN).save(SAMPLES / "karan_mehta_ios_photo.png")
    text_pdf(SAMPLES / "anita_desai_accountant.pdf", ANITA, {})
    (SAMPLES / "pierre_martin_android_fr.txt").write_text(PIERRE, encoding="utf-8")
    docx_file(SAMPLES / "sneha_kulkarni_android_hi.docx", SNEHA)
    docx_file(SAMPLES / "rahul_verma_web.docx", RAHUL, ("https://github.com/octocat", "GitHub"))
    for path in sorted(SAMPLES.iterdir()):
        print(f"{path.stat().st_size:>8,} B  {path.name}")


if __name__ == "__main__":
    main()
