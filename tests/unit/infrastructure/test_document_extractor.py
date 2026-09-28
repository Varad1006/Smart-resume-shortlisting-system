import io
import shutil
from concurrent.futures import ThreadPoolExecutor

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fpdf import FPDF
from PIL import Image, ImageDraw, ImageFont

from resume_shortlister.domain.errors import (
    ExtractionError,
    OcrUnavailableError,
    UnsupportedDocumentError,
)
from resume_shortlister.domain.models import ExtractionMethod
from resume_shortlister.infrastructure.extraction.document_extractor import (
    DocumentExtractor,
    DocumentKind,
    detect_kind,
    normalize_text,
)
from resume_shortlister.infrastructure.ocr.base import OcrOutput
from resume_shortlister.infrastructure.ocr.tesseract import TesseractOcrEngine

TEXT_LINE = "Senior Android developer with five years of Kotlin and Jetpack Compose."


class FakeOcr:
    def __init__(self, fail: bool = False) -> None:
        self.fail = fail
        self.batches: list[int] = []

    def is_available(self) -> bool:
        return True

    def status(self):
        return []

    def recognize(self, images):
        self.batches.append(len(images))
        if self.fail:
            raise ExtractionError("OCR server exploded")
        return OcrOutput(
            texts=[f"Scanned page text number {i}" for i in range(len(images))],
            method=ExtractionMethod.CHANDRA,
        )


def text_image(text: str = "HELLO RESUME", size=(900, 300)) -> Image.Image:
    image = Image.new("RGB", size, "white")
    ImageDraw.Draw(image).text((40, 100), text, fill="black", font=ImageFont.load_default(size=64))
    return image


def make_pdf(pages: list[str | Image.Image], link: str | None = None) -> bytes:
    pdf = FPDF()
    pdf.set_font("Helvetica", size=12)
    for page in pages:
        pdf.add_page()
        if isinstance(page, str):
            pdf.multi_cell(0, 8, page)
            if link:
                pdf.cell(40, 10, "GitHub", link=link)
        else:
            pdf.image(page, x=10, y=10, w=180)
    return bytes(pdf.output())


def png_bytes(image: Image.Image) -> bytes:
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def make_docx() -> bytes:
    document = Document()
    header = document.sections[0].header
    header.is_linked_to_previous = False
    header.paragraphs[0].text = "Jane Doe · jane@example.com"
    paragraph = document.add_paragraph("Profile: ")
    rel_id = paragraph.part.relate_to(
        "https://github.com/jane-doe", RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "my GitHub"
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    document.add_paragraph(TEXT_LINE)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "Kotlin, Firebase"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_pdf_text_layer_and_hidden_links():
    ocr = FakeOcr()
    doc = DocumentExtractor(ocr).extract("cv.pdf", make_pdf([TEXT_LINE], "https://github.com/jane"))
    assert doc.methods == (ExtractionMethod.TEXT_LAYER,)
    assert "Jetpack Compose" in doc.text
    assert doc.links == ("https://github.com/jane",)
    assert ocr.batches == []  # no OCR needed


def test_scanned_pdf_goes_through_ocr():
    ocr = FakeOcr()
    doc = DocumentExtractor(ocr).extract("scan.pdf", make_pdf([text_image(), text_image()]))
    assert doc.methods == (ExtractionMethod.CHANDRA,)
    assert ocr.batches == [2]  # both pages in one batch
    assert doc.pages == 2


def test_mixed_pdf_only_ocrs_scanned_pages():
    ocr = FakeOcr()
    doc = DocumentExtractor(ocr).extract("mixed.pdf", make_pdf([TEXT_LINE, text_image()]))
    assert doc.methods == (ExtractionMethod.TEXT_LAYER, ExtractionMethod.CHANDRA)
    assert ocr.batches == [1]
    assert "Kotlin" in doc.text and "Scanned page text" in doc.text


def test_mixed_pdf_keeps_text_layer_when_ocr_fails():
    doc = DocumentExtractor(FakeOcr(fail=True)).extract(
        "m.pdf", make_pdf([TEXT_LINE, text_image()])
    )
    assert doc.methods == (ExtractionMethod.TEXT_LAYER,)
    assert "could not be read" in doc.warnings[0]


def test_scanned_pdf_without_ocr_is_rejected():
    with pytest.raises(OcrUnavailableError):
        DocumentExtractor(None).extract("scan.pdf", make_pdf([text_image()]))


def test_page_limit_adds_a_warning():
    doc = DocumentExtractor(None, max_pages=2).extract("long.pdf", make_pdf([TEXT_LINE] * 3))
    assert doc.pages == 2
    assert doc.warnings == ("Only the first 2 of 3 pages were read.",)


def test_concurrent_pdf_extraction_is_thread_safe():
    # Regression: PDFium is not thread-safe; parallel extraction used to abort the process.
    extractor = DocumentExtractor(FakeOcr())
    documents = [make_pdf([f"{TEXT_LINE} Candidate {i}.", text_image()]) for i in range(16)]
    with ThreadPoolExecutor(max_workers=8) as pool:
        results = list(pool.map(lambda pdf: extractor.extract("cv.pdf", pdf), documents * 3))
    assert len(results) == 48
    assert all(
        r.methods == (ExtractionMethod.TEXT_LAYER, ExtractionMethod.CHANDRA) for r in results
    )


def test_corrupt_pdf_raises_extraction_error():
    with pytest.raises(ExtractionError, match="Could not open this PDF"):
        DocumentExtractor(None).extract("bad.pdf", b"%PDF-1.7 garbage")


def test_docx_text_header_table_and_hyperlinks():
    doc = DocumentExtractor(None).extract("cv.docx", make_docx())
    assert doc.methods == (ExtractionMethod.DOCX,)
    assert doc.text.splitlines()[0] == "Jane Doe · jane@example.com"
    assert "Profile: my GitHub" in doc.text
    assert "Kotlin, Firebase" in doc.text
    assert doc.links == ("https://github.com/jane-doe",)


def test_image_is_ocred_including_multi_page_tiff():
    ocr = FakeOcr()
    extractor = DocumentExtractor(ocr)
    assert extractor.extract("photo.png", png_bytes(text_image())).methods == (
        ExtractionMethod.CHANDRA,
    )
    buffer = io.BytesIO()
    text_image().save(buffer, format="TIFF", save_all=True, append_images=[text_image()])
    assert extractor.extract("scan.tiff", buffer.getvalue()).pages == 2
    assert ocr.batches == [1, 2]


def test_image_without_ocr_is_rejected():
    with pytest.raises(OcrUnavailableError):
        DocumentExtractor(None).extract("photo.png", png_bytes(text_image()))


def test_plain_text():
    doc = DocumentExtractor(None).extract("cv.txt", TEXT_LINE.encode())
    assert doc.methods == (ExtractionMethod.PLAIN_TEXT,)
    assert doc.text == TEXT_LINE


def test_unknown_binary_is_unsupported():
    with pytest.raises(UnsupportedDocumentError):
        DocumentExtractor(None).extract("cv.pdf", bytes(range(256)) * 4)


@pytest.mark.parametrize(
    ("content", "kind"),
    [
        (make_pdf([TEXT_LINE]), DocumentKind.PDF),
        (make_docx(), DocumentKind.DOCX),
        (png_bytes(text_image()), DocumentKind.IMAGE),
        ("नमस्ते resume".encode(), DocumentKind.TEXT),
    ],
)
def test_detect_kind_uses_content_not_extension(content, kind):
    assert detect_kind(content) is kind


def test_normalize_text():
    raw = "Proﬁcient\r\n\r\n\r\n\r\nPython   developer\x00\n  \n"
    assert normalize_text(raw) == "Proficient\n\nPython developer"


@pytest.mark.skipif(shutil.which("tesseract") is None, reason="tesseract binary not installed")
def test_tesseract_reads_rendered_text():
    engine = TesseractOcrEngine("eng")
    assert engine.is_available()
    output = engine.recognize([text_image("KOTLIN DEVELOPER")])
    assert output.method is ExtractionMethod.TESSERACT
    assert "KOTLIN" in output.texts[0].upper()
